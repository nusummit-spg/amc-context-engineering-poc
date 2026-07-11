"""WS2 — PPTX parser (python-pptx): one section per slide + speaker notes."""
from pathlib import Path

from pptx import Presentation

from app.ingestion.parsers.base import BaseParser
from app.schemas.documents import Document, DocumentType, Section


class PptxParser(BaseParser):
    extensions = (".pptx",)

    def parse(self, path: Path) -> Document:
        prs = Presentation(str(path))
        sections: list[Section] = []

        for index, slide in enumerate(prs.slides, start=1):
            lines: list[str] = []
            title = None
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                text = shape.text_frame.text.strip()
                if not text:
                    continue
                if title is None and shape == slide.shapes.title:
                    title = text
                else:
                    lines.append(text)
            notes = ""
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
            body = "\n".join(lines)
            if notes:
                body += f"\n\n[SPEAKER NOTES]\n{notes}"
            if title or body.strip():
                sections.append(Section(
                    title=title or f"Slide {index}",
                    text=body.strip() or (title or ""),
                    order=index - 1,
                    metadata={"slide": index},
                ))

        return Document(
            filename=path.name,
            doc_type=DocumentType.PPTX,
            title=path.stem.replace("_", " "),
            source_path=str(path),
            sections=sections or [Section(text="", order=0)],
            metadata={"slide_count": len(sections)},
        )
