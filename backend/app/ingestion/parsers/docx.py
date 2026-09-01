# ===========================================================================
# Copyright © NuSummit Technologies Pvt Ltd. All rights reserved.
# 
# Author: NuSummit Developers
#
# ===========================================================================

"""WS2 — DOCX parser (python-docx): heading hierarchy preserved for chunking."""
from pathlib import Path

from docx import Document as DocxFile

from app.ingestion.parsers.base import BaseParser
from app.schemas.documents import Document, DocumentType, Section


class DocxParser(BaseParser):
    extensions = (".docx",)

    def parse(self, path: Path) -> Document:
        docx = DocxFile(str(path))
        sections: list[Section] = []
        current_title: str | None = None
        current_level = 1
        current_lines: list[str] = []
        order = 0

        def flush():
            nonlocal order
            text = "\n".join(current_lines).strip()
            if text:
                sections.append(Section(
                    title=current_title, level=current_level, text=text, order=order,
                ))
                order += 1

        for para in docx.paragraphs:
            style = (para.style.name or "").lower()
            if style.startswith("heading"):
                flush()
                current_title = para.text.strip()
                try:
                    current_level = int(style.replace("heading", "").strip() or 1)
                except ValueError:
                    current_level = 1
                current_lines = []
            elif para.text.strip():
                current_lines.append(para.text)

        # Tables as pipe-delimited rows appended to the current section.
        for table in docx.tables:
            rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
            if rows:
                current_lines.append("\n[TABLE]\n" + "\n".join(rows))
        flush()

        author = docx.core_properties.author or None
        return Document(
            filename=path.name,
            doc_type=DocumentType.DOCX,
            title=docx.core_properties.title or path.stem.replace("_", " "),
            author=author,
            source_path=str(path),
            sections=sections or [Section(text="", order=0)],
        )
